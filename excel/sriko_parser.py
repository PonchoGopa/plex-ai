"""
excel/sriko_parser.py

Parser determinístico para Release Schedule de S-Riko (Excel .xlsx).

Reglas de negocio:
  - PO No fijo: extraído de la celda T12
  - Customer fijo: extraído de la celda I8
  - Solo se procesan semanas FIRM
  - Filas con cantidad 0 en una semana se ignoran para esa semana
  - Partes cuya cantidad es 0 en TODAS las semanas FIRM se ignoran
"""
from __future__ import annotations

import logging
from datetime import datetime

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from pdf.mos_table_parser import MosHeader, MosRecord

logger = logging.getLogger(__name__)

_ROW_CUSTOMER   = 8
_COL_CUSTOMER   = 9    # columna I

_ROW_PO         = 12
_COL_PO         = 20   # columna T

_ROW_FIRM_FLAG  = 16
_ROW_DATES      = 18
_ROW_DATA_START = 19

_COL_PART_NO    = 2    # columna B
_COL_SNP        = 4    # columna D
_COL_DATA_START = 6    # columna F — primera semana


class SRikoParser:

    def parse(self, file_path: str) -> tuple[MosHeader, list[MosRecord]]:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active

        header    = self._parse_header(ws)
        firm_cols = self._detect_firm_columns(ws)
        records   = self._parse_records(ws, firm_cols)

        logger.info(
            "SRikoParser → customer=%r | po=%r | %d semanas FIRM | %d registros",
            header.customer,
            header.po_number,
            len(firm_cols),
            len(records),
        )
        return header, records

    def _parse_header(self, ws: Worksheet) -> MosHeader:
        customer = self._cell_str(ws, _ROW_CUSTOMER, _COL_CUSTOMER)
        po_raw   = self._cell_str(ws, _ROW_PO, _COL_PO)
        po_number = po_raw.split()[0] if po_raw else ""
        return MosHeader(customer=customer or "S-Riko", po_number=po_number)

    def _detect_firm_columns(self, ws: Worksheet) -> list[int]:
        firm_cols: list[int] = []
        max_col = ws.max_column or 30
        in_firm = True

        for col in range(_COL_DATA_START, max_col + 1):
            val = self._cell_str(ws, _ROW_FIRM_FLAG, col).upper()
            if val and ("FORE" in val or "FORECAST" in val):
                in_firm = False
            if in_firm:
                date_val = ws.cell(_ROW_DATES, col).value
                if isinstance(date_val, datetime):
                    firm_cols.append(col)

        return firm_cols

    def _parse_records(
        self,
        ws: Worksheet,
        firm_cols: list[int],
    ) -> list[MosRecord]:
        records: list[MosRecord] = []

        # Fechas de columnas FIRM
        col_dates: dict[int, datetime] = {}
        for col in firm_cols:
            val = ws.cell(_ROW_DATES, col).value
            if isinstance(val, datetime):
                col_dates[col] = val

        if not col_dates:
            logger.warning("SRikoParser: no se encontraron fechas FIRM.")
            return []

        row = _ROW_DATA_START
        max_row = ws.max_row or 200
        empty_streak = 0

        while row <= max_row:
            part_no = self._cell_str(ws, row, _COL_PART_NO)

            if not part_no:
                empty_streak += 1
                if empty_streak >= 3:
                    break
                row += 1
                continue

            empty_streak = 0
            snp = self._cell_str(ws, row, _COL_SNP)

            for col, fecha in col_dates.items():
                qty_val = ws.cell(row, col).value
                qty     = self._parse_qty(qty_val)

                if qty == 0:
                    continue

                records.append(MosRecord(
                    part_number  = part_no,
                    part_name    = "",              # S-Riko no tiene part name
                    model        = "",              # S-Riko no tiene model
                    snp          = snp,
                    date         = fecha.strftime("%d/%m/%Y"),
                    quantity_box = 0,
                    quantity_qty = qty,
                    po_number    = None,            # tomar del header
                ))

            row += 1

        return records

    @staticmethod
    def _cell_str(ws: Worksheet, row: int, col: int) -> str:
        val = ws.cell(row, col).value
        return str(val).strip() if val is not None else ""

    @staticmethod
    def _parse_qty(val) -> int:
        if val is None:
            return 0
        try:
            return int(float(str(val).replace(",", "").strip()))
        except (ValueError, AttributeError):
            return 0